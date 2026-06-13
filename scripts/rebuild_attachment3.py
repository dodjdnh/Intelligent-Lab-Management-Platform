from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


path = Path(r"C:\Users\Kevin\Desktop\智能实验室管理项目\附件资料\附件3.上海海洋大学大学生创新大赛（2026）报名表汇总表(XX学院姓名学号).xlsx")

headers = [
    "序号",
    "项目名称",
    "项目简介\n（100-200字）",
    "主要创新点和难点\n（100字以内）",
    "项目组别\n（主赛道本科生组/\n主赛道研究生组/\n红旅组）",
    "项目类别\n（主赛道本、研：创意组/创业组；\n红旅赛道：公益组/创意组/创业组）",
    "项目类型\n（仅主赛道填写：新工科/新医科/新农科/新文科/人工智能+/低空经济/生物技术/量子科技/新能源/新材料）",
    "负责人",
    "负责人学号",
    "负责人所在学院",
    "负责人电话",
    "负责人邮箱",
    "参与学生人数\n(包含负责人)",
    "项目其他成员信息\n（不包含负责人）",
    "指导教师",
    "职称",
    "研究方向",
    "教师电话",
    "教师邮箱",
    "教师所在学院",
    "备注",
]

row2 = [
    1,
    "智衡实验云枢——实验室智慧管理平台",
    "本项目面向高校实验室数字化管理场景，围绕预约审批、耗材库存、设备监测、文件中心与智能助手构建统一平台。系统结合 Spring Boot、Vue、MySQL 与物联网感知能力，实现实验资源调度、库存追踪、设备状态展示和异常告警联动，提升实验室运行效率、透明度与安全性，并具备向智能运营平台持续扩展的基础。",
    "创新点在于将实验室预约、耗材管理、设备接入和 AI/IoT 感知能力统一到同一平台，实现库存与设备状态实时联动；难点在于多角色流程协同、传感器数据与业务规则映射，以及系统实时性与可扩展性的平衡。",
    "主赛道本科生组",
    "创意组",
    "人工智能+",
    "姚佳昂",
    "2551226",
    "信息学院",
    "18837236586",
    "3478917829@qq.com",
    "5",
    "雷璐灵/2550212,\n鄺健豪/2550330,\n许家怿/2534315,\n欧阳中昊/2551217",
    "赵慧娟/待补充，张增敏/待补充",
    "讲师，教师",
    "算法设计；工程数据库、机器人工程、人工智能、程序设计",
    "15692165685，15618063223",
    "待补充",
    "信息学院，工程学院",
    "立项大创项目，需同步完成官网和学院报名",
]

note = (
    "注意：1.团队学生信息、指导老师信息请务必填写正确且完整的信息，并确保排序正确"
    "（获奖团队将按此表格信息进行申报，提交后不可更改）。\n"
    "2.核对团队学生信息、指导教师信息、报名赛道、组别、项目类别、类型等信息，务必和报名表一致。"
)

wb = Workbook()
ws = wb.active
ws.title = "报名汇总表"

for idx, header in enumerate(headers, start=1):
    ws.cell(1, idx).value = header
for idx, value in enumerate(row2, start=1):
    ws.cell(2, idx).value = value

ws.cell(3, 1).value = 2
ws.cell(4, 1).value = note
ws.merge_cells(start_row=4, start_column=1, end_row=4, end_column=len(headers))

fill = PatternFill("solid", fgColor="1A3C5E")
white_font = Font(color="FFFFFF", bold=True, name="Microsoft YaHei")
body_font = Font(name="Microsoft YaHei", size=10)
thin = Side(style="thin", color="B8C7D9")
border = Border(left=thin, right=thin, top=thin, bottom=thin)

for row in ws.iter_rows(min_row=1, max_row=4, min_col=1, max_col=len(headers)):
    for cell in row:
        cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
        cell.border = border
        cell.font = body_font

for cell in ws[1]:
    cell.fill = fill
    cell.font = white_font

ws["B2"].alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
for col in ("C", "D", "N", "Q", "U"):
    ws[f"{col}2"].alignment = Alignment(vertical="top", horizontal="left", wrap_text=True)

ws["A4"].alignment = Alignment(vertical="center", horizontal="left", wrap_text=True)

widths = {
    "A": 8, "B": 24, "C": 34, "D": 28, "E": 16, "F": 18, "G": 20,
    "H": 10, "I": 12, "J": 14, "K": 14, "L": 22, "M": 12, "N": 22,
    "O": 18, "P": 12, "Q": 26, "R": 16, "S": 16, "T": 18, "U": 20,
}
for col, width in widths.items():
    ws.column_dimensions[col].width = width

ws.row_dimensions[1].height = 72
ws.row_dimensions[2].height = 84
ws.row_dimensions[4].height = 42

wb.save(path)
print(path)
